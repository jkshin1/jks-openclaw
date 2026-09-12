#!/usr/bin/env python3
"""Synthetic fixtures for acceptance of the user's native OpenClaw office runtime.

This is an integration test of the specifically installed python-docx/openpyxl
tools, not an end-user document authoring workflow. No owner data is read.
"""

import argparse
from datetime import datetime
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from pypdf import PdfReader


def build(root):
    root.mkdir(mode=0o700, parents=True, exist_ok=False)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    for name in ["Normal", "Title", "Heading 1"]:
        style = doc.styles[name]
        style.font.name = "Apple SD Gothic Neo"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Apple SD Gothic Neo")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph("문서 기능 검증", "Title")
    doc.add_paragraph("한국어 문장과 표를 생성하고 지정한 내용만 수정하는 시험입니다.")
    p = doc.add_paragraph("출시 상태: ")
    p.add_run("검토 중").bold = True
    doc.add_paragraph("유지할 문장: 자료의 출처와 원본 파일을 보존합니다.")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Inches(1.9), Inches(4.7)
    for cell, text in zip(table.rows[0].cells, ["항목", "확인 내용"]):
        cell.text = text
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for left, right in [("한국어", "문장과 숫자 12345"), ("수정 범위", "굵은 상태 문구 한 곳"),
                        ("표 보존", "열 너비와 행 내용을 유지")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = left, right
    section.header.paragraphs[0].text = "OpenClaw 문서 시험"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_page_break()
    doc.add_paragraph("두 번째 페이지", "Heading 1")
    doc.add_paragraph("페이지 구분과 머리글 및 바닥글이 유지되는지 확인합니다.")
    doc.save(root / "document-input.docx")

    book = Workbook()
    sheet = book.active
    sheet.title = "작업"
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "작업 비용 검증"
    sheet["A1"].font = Font(name="Apple SD Gothic Neo", size=18, bold=True)
    sheet.append([])
    sheet.append(["업무 ID", "작업", "수량", "단가", "금액", "완료일", "상태"])
    rows = [("0007", "문서 정리", 3, 12000), ("001234567890123456", "표 점검", 2, 18000),
            ("0010", "음성 정리", 5, 7000)]
    for row, (identifier, label, count, price) in enumerate(rows, 4):
        for col, value in enumerate([identifier, label, count, price, f"=C{row}*D{row}",
                                     datetime(2026, 9, 7), "대기"], 1):
            sheet.cell(row, col, value)
        sheet.cell(row, 1).number_format = "@"
        sheet.cell(row, 6).number_format = "yyyy-mm-dd"
        for col in (4, 5):
            sheet.cell(row, col).number_format = '#,##0'
    for row, label, formula in [(8, "소계", "=SUM(E4:E6)"),
                                (9, "추가 금액", "=E8*'설정'!$B$3"),
                                (10, "합계", "=SUM(E8:E9)")]:
        sheet.cell(row, 4, label)
        sheet.cell(row, 5, formula).number_format = '#,##0'
    settings = book.create_sheet("설정")
    settings.append(["설정", "값"])
    settings.append(["용도", "합성 수식 시험"])
    settings.append(["추가 비율", 0.1])
    settings["B3"].number_format = "0.0%"
    book.defined_names.add(DefinedName("ExtraRate", attr_text="'설정'!$B$3"))
    hidden = book.create_sheet("숨김")
    hidden["A1"] = "KEEP-001"
    hidden.sheet_state = "hidden"
    validation = DataValidation(type="list", formula1='"대기,완료"')
    sheet.add_data_validation(validation)
    validation.add("G4:G6")
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = "A3:G6"
    widths = {"A": 25, "B": 16, "C": 8, "D": 13, "E": 15, "F": 16, "G": 10}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=3, max_row=10, max_col=7):
        for cell in row:
            cell.font = Font(name="Apple SD Gothic Neo", size=11)
            cell.alignment = Alignment(vertical="center")
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Apple SD Gothic Neo", color="FFFFFF", bold=True, size=11)
    for row in range(3, 11):
        sheet.row_dimensions[row].height = 25
    sheet.print_area = "A1:G11"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    settings.column_dimensions["A"].width = 20
    settings.column_dimensions["B"].width = 28
    settings.print_area = "A1:B3"
    book.save(root / "workbook-input.xlsx")

    (root / "summary-input.txt").write_text(
        "가상의 도서관 운영 회의 기록입니다. 실제 사용자나 조직의 개인정보는 없습니다.\n"
        "도서관은 다음 달부터 토요일 운영을 두 시간 연장하기로 확정했습니다. 평일 운영 시간은 그대로 유지합니다.\n"
        "새로운 예약 시스템을 도입하기 전에 직원 네 명이 일주일 동안 시험 운영을 진행합니다. "
        "시험에서 발견한 문제는 기록하고 수정한 뒤 전체 이용자에게 공개할 예정입니다.\n"
        "이번에 승인된 자료 정리 예산은 삼백만 원입니다. 추가 장비 구매는 아직 결정하지 않았습니다. "
        "시험 운영 결과를 보고 다음 회의에서 장비 필요 여부를 판단합니다.\n"
        "홍보물은 한국어로 먼저 작성합니다. 운영 안내에는 바뀐 토요일 시간과 예약 방법을 함께 적습니다. "
        "현재 홈페이지의 잘못된 연락처도 수정하기로 했습니다.\n"
        "확정된 결정과 아직 결정하지 않은 사항을 구분해 요약해 주세요.\n", encoding="utf-8")
    (root / "speech-source.txt").write_text(
        "안녕하세요. 오늘 회의에서는 문서 정리와 일정 변경을 확인했습니다. "
        "새로운 보고서는 금요일까지 작성합니다. 중요한 내용은 한국어로 간단하게 정리해 주세요.", encoding="utf-8")
    print(json.dumps({"fixtures": str(root)}, ensure_ascii=False))


def check_office(root):
    original = Document(root / "document-input.docx")
    assert any(run.text == "검토 중" and run.bold for p in original.paragraphs for run in p.runs)
    assert load_workbook(root / "workbook-input.xlsx")["작업"]["C5"].value == 2
    doc = Document(root / "document-output.docx")
    assert any(run.text == "검증 완료" and run.bold for p in doc.paragraphs for run in p.runs)
    assert any(p.text == "유지할 문장: 자료의 출처와 원본 파일을 보존합니다." for p in doc.paragraphs)
    assert len(doc.tables) == 1 and doc.tables[0].cell(3, 1).text == "열 너비와 행 내용을 유지"
    assert doc.sections[0].header.paragraphs[0].text == "OpenClaw 문서 시험"
    path = root / "recalculated/workbook-output.xlsx"
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    sheet = formulas["작업"]
    assert sheet["A4"].value == "0007" and sheet["A5"].value == "001234567890123456"
    assert sheet["C5"].value == 4 and sheet["E5"].data_type == "f"
    assert values["작업"]["E10"].value == 157300
    assert sheet.freeze_panes == "A4" and sheet.auto_filter.ref == "A3:G6"
    assert list(sheet.data_validations.dataValidation) and "ExtraRate" in formulas.defined_names
    assert formulas["숨김"].sheet_state == "hidden" and formulas["숨김"]["A1"].value == "KEEP-001"
    for tab in values:
        for row in tab:
            assert not any(cell.data_type == "e" for cell in row), "formula error"
    word_pdf = PdfReader(root / "document-output.pdf")
    excel_pdf = PdfReader(root / "recalculated/workbook-output.pdf")
    assert len(word_pdf.pages) == len(excel_pdf.pages) == 2
    word_text = "\n".join(page.extract_text() for page in word_pdf.pages)
    excel_text = "\n".join(page.extract_text() for page in excel_pdf.pages)
    assert "검증 완료" in word_text and "두 번째 페이지" in word_text, "Korean Word PDF text lost"
    assert "작업 비용 검증" in excel_text and "157,300" in excel_text, "Korean Excel PDF text lost"
    print(json.dumps({"wordTargetedEdit": True, "wordPreservedStructure": True,
                      "excelRecalculatedTotal": 157300, "excelPreservedStructureAndTypes": True,
                      "originalInputsUnedited": True, "koreanPdfText": True, "pdfPagesEach": 2}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=["build", "check-office"])
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    (build if args.mode == "build" else check_office)(args.directory)
