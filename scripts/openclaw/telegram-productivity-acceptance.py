#!/usr/bin/env python3
"""Stage, verify and deliberately deliver synthetic Office/media acceptance files.

Run prepare/verify with openclaw-python to exercise the installed native engines.
No model/API fallback, owner document discovery or Telegram polling is performed.
The receipt distinguishes native processing, reviewed rendering and outbound
Telegram acceptance from an actual owner upload or viewing files on a phone.
"""

import argparse
from contextlib import redirect_stdout
from datetime import datetime, timezone
import fcntl
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import re
import signal
import shutil
import subprocess
import sys
import tempfile
import zipfile


STATE = Path.home() / ".openclaw-personaledge"
CLI = Path.home() / ".local/openclaw-2026.8.1/.personal-edge-management/bin/openclaw"
INPUTS = ("document-input.docx", "workbook-input.xlsx", "speech-source.txt")
ARTIFACTS = (
    ("docx", "document-output.docx"),
    ("xlsx", "recalculated/workbook-output.xlsx"),
    ("pdf-word", "document-output.pdf"),
    ("pdf-excel", "recalculated/workbook-output.pdf"),
    ("audio", "speech-input.wav"),
    ("transcript", "transcripts/speech-input.txt"),
    ("subtitles", "transcripts/speech-input.srt"),
)


def require(condition, message):
    if not condition:
        raise ValueError(message)


def now():
    return datetime.now(timezone.utc).isoformat()


def sha(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_file(root, relative):
    path = root / relative
    require(not Path(relative).is_absolute() and ".." not in Path(relative).parts,
            "unsafe artifact path")
    require(path.is_file() and not any(p.is_symlink() for p in [path, *path.parents]),
            "missing or redirected artifact: " + relative)
    require(path.resolve().is_relative_to(root), "artifact escaped run directory")
    return path


def save(path, payload):
    require(not path.is_symlink(), "receipt path is a symlink")
    fd, name = tempfile.mkstemp(prefix=".receipt-", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, path)
        directory = os.open(path.parent, os.O_RDONLY)
        try:
            os.fsync(directory)
        finally:
            os.close(directory)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def read_receipt(root, name):
    return json.loads(safe_file(root, name).read_text(encoding="utf-8"))


def bounded(command, timeout=180, env=None, diagnostic_path=None):
    """Bound native process groups, including children of Office/Whisper wrappers."""
    started = now()
    if diagnostic_path is not None:
        require(not any(path.is_symlink() for path in [diagnostic_path, *diagnostic_path.parents]),
                "diagnostic path is redirected")
        diagnostic_path.parent.mkdir(mode=0o700, parents=True, exist_ok=True)

    def diagnose(stdout, stderr, code, timed_out=False):
        if diagnostic_path is not None:
            save(diagnostic_path, {"command": Path(str(command[0])).name,
                                  "startedAt": started, "completedAt": now(),
                                  "exitCode": code, "timedOut": timed_out,
                                  "stdout": stdout[-131072:], "stderr": stderr[-131072:],
                                  "truncated": len(stdout) > 131072 or len(stderr) > 131072})

    process = subprocess.Popen([str(arg) for arg in command], stdout=subprocess.PIPE,
                               stderr=subprocess.PIPE, text=True, start_new_session=True,
                               env=env)
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        stdout, stderr = process.communicate()
        diagnose(stdout, stderr, process.returncode, timed_out=True)
        raise ValueError(Path(str(command[0])).name + " timed out") from None
    diagnose(stdout, stderr, process.returncode)
    require(process.returncode == 0, Path(str(command[0])).name +
            " failed (exit " + str(process.returncode) + "); " +
            ("private diagnostics saved" if diagnostic_path else "no raw diagnostics echoed"))
    return stdout


def fixture_module():
    spec = importlib.util.spec_from_file_location("productivity_fixtures",
                                                 Path(__file__).with_name("productivity-fixtures.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def normalized(text):
    return "".join(char for char in text if char.isalnum())


def artifact_hashes(root):
    return {name: sha(safe_file(root, name)) for _, name in ARTIFACTS}


def render_hashes(root):
    names = ["render/word-page-1.png", "render/word-page-2.png",
             "render/excel-page-1.png", "render/excel-page-2.png"]
    return {name: sha(safe_file(root, name)) for name in names}


def subtitle_archive(root):
    """Package exact SRT bytes as a validated ZIP without relaxing media checks."""
    source = safe_file(root, "transcripts/speech-input.srt")
    payload = source.read_bytes()
    archive_bytes = io.BytesIO()
    info = zipfile.ZipInfo("speech-input.srt", date_time=(1980, 1, 1, 0, 0, 0))
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o600 << 16
    with zipfile.ZipFile(archive_bytes, "w") as archive:
        archive.writestr(info, payload)
    expected = archive_bytes.getvalue()
    name = "transcripts/speech-input-subtitles.zip"
    target = root / name
    require(not target.is_symlink(), "subtitle archive is redirected")
    if not target.exists():
        fd = os.open(target, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "wb") as stream:
            stream.write(expected)
            stream.flush()
            os.fsync(stream.fileno())
    require(safe_file(root, name).read_bytes() == expected, "subtitle archive content changed")
    with zipfile.ZipFile(target) as archive:
        require(archive.namelist() == ["speech-input.srt"] and archive.testzip() is None and
                archive.read("speech-input.srt") == payload, "subtitle archive entry mismatch")
    return {"format": "zip", "path": name, "sha256": sha(target),
            "entry": "speech-input.srt", "entrySha256": sha(source)}


def prepare(args):
    from docx import Document
    from openpyxl import load_workbook
    import pypdfium2 as pdfium

    root = args.run_dir
    require(not root.exists(), "prepare requires a fresh run directory")
    # Fail before generating files rather than let Whisper download a model.
    model = Path.home() / ".local/share/openclaw-skill-tools/models/whisper/small.pt"
    require(model.is_file(), "installed local Whisper small model missing; no download attempted")
    fixture = fixture_module()
    with redirect_stdout(io.StringIO()):
        fixture.build(root)
    receipt = {"schemaVersion": 1, "createdAt": now(), "syntheticOnly": True,
               "processingEngine": "installed-native-openclaw-tools",
               "llmInference": False, "localWhisperInference": True,
               "inputs": {name: sha(root / name) for name in INPUTS},
               "ownerUploadVerified": False, "phoneOpenVerified": False}
    save(root / "manifest.json", receipt)
    document = Document(root / "document-input.docx")
    edits = 0
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text == "검토 중" and run.bold:
                run.text = "검증 완료"
                edits += 1
    require(edits == 1, "unexpected DOCX target count")
    document.save(root / "document-output.docx")
    book = load_workbook(root / "workbook-input.xlsx")
    book["작업"]["C5"] = 4
    book.save(root / "workbook-output.xlsx")
    (root / "recalculated").mkdir(mode=0o700)
    office = [args.office, "-env:UserInstallation=" + (root / "office-profile").as_uri(),
              "--headless"]
    bounded([*office, "--convert-to", "pdf", "--outdir", root, root / "document-output.docx"])
    bounded([*office, "--convert-to", "xlsx", "--outdir", root / "recalculated",
             root / "workbook-output.xlsx"])
    bounded([*office, "--convert-to", "pdf", "--outdir", root / "recalculated",
             root / "recalculated/workbook-output.xlsx"])
    (root / "render").mkdir(mode=0o700)
    for label, source in (("word", "document-output.pdf"),
                          ("excel", "recalculated/workbook-output.pdf")):
        pdf = pdfium.PdfDocument(root / source)
        try:
            require(len(pdf) == 2, label + " PDF must contain two pages")
            for index in range(len(pdf)):
                page = pdf[index]
                bitmap = page.render(scale=1.5)
                bitmap.to_pil().save(root / "render" / f"{label}-page-{index + 1}.png")
                bitmap.close()
                page.close()
        finally:
            pdf.close()
    bounded(["/usr/bin/say", "-v", "Yuna", "-f", root / "speech-source.txt",
             "-o", root / "speech-input.aiff"])
    bounded([args.ffmpeg, "-nostdin", "-v", "error", "-i", root / "speech-input.aiff",
             "-ac", "1", "-ar", "16000", root / "speech-input.wav"])
    (root / "transcripts").mkdir(mode=0o700)
    bounded([args.whisper, root / "speech-input.wav", "--language", "Korean",
             "--output_format", "all", "--output_dir", root / "transcripts"], timeout=300)
    receipt["artifacts"] = artifact_hashes(root)
    receipt["renderedPages"] = render_hashes(root)
    receipt["transportArtifacts"] = {"transcripts/speech-input.srt": subtitle_archive(root)}
    save(root / "manifest.json", receipt)
    return verify(args)


def verify(args):
    root = args.run_dir
    manifest = read_receipt(root, "manifest.json")
    require(manifest.get("schemaVersion") == 1 and manifest.get("syntheticOnly") is True,
            "unrecognized synthetic manifest")
    require(set(manifest.get("inputs", {})) == set(INPUTS), "incomplete original input hashes")
    for name, expected in manifest["inputs"].items():
        require(sha(safe_file(root, name)) == expected, "original input changed: " + name)
    require(manifest.get("artifacts") == artifact_hashes(root), "output artifact changed")
    require(manifest.get("renderedPages") == render_hashes(root), "rendered page changed")
    with redirect_stdout(io.StringIO()) as captured:
        fixture_module().check_office(root)
    checks = json.loads(captured.getvalue())
    source = safe_file(root, "speech-source.txt").read_text(encoding="utf-8")
    transcript = safe_file(root, "transcripts/speech-input.txt").read_text(encoding="utf-8")
    require(normalized(source) == normalized(transcript), "Korean synthetic transcript mismatch")
    speech = read_receipt(root, "transcripts/speech-input.json")
    require(speech.get("language") in ("ko", "Korean") and speech.get("segments"),
            "Whisper transcript segments/language missing")
    last_end = 0
    for segment in speech["segments"]:
        require(isinstance(segment.get("start"), (int, float)) and
                isinstance(segment.get("end"), (int, float)) and
                0 <= segment["start"] <= segment["end"] and segment["start"] >= last_end,
                "invalid Whisper segment timing")
        last_end = segment["end"]
    require(normalized(speech.get("text", "")) == normalized(source), "JSON transcript mismatch")
    srt = safe_file(root, "transcripts/speech-input.srt").read_text(encoding="utf-8")
    require(re.search(r"\d\d:\d\d:\d\d,\d{3} --> \d\d:\d\d:\d\d,\d{3}", srt),
            "SRT subtitle timing missing")
    subtitle_text = " ".join(line for line in srt.splitlines()
                             if line.strip() and not line.isdigit() and " --> " not in line)
    require(normalized(subtitle_text) == normalized(source), "SRT transcript mismatch")
    for name in ("speech-input.vtt", "speech-input.tsv"):
        require(safe_file(root, "transcripts/" + name).stat().st_size > 0,
                "empty subtitle companion")
    transport_artifacts = {"transcripts/speech-input.srt": subtitle_archive(root)}
    if "transportArtifacts" in manifest:
        require(manifest["transportArtifacts"] == transport_artifacts, "subtitle transport hash changed")
    checks.update({"originalInputHashesPreserved": True, "syntheticKoreanTranscript": True,
                   "subtitleTiming": True, "renderedPages": 4})
    result = {"schemaVersion": 1, "verifiedAt": now(), "nativeProcessing": "passed",
              "checks": checks, "artifacts": manifest["artifacts"],
              "transportArtifacts": transport_artifacts,
              "renderedPages": manifest["renderedPages"], "visualReview": "pending",
              "ownerUploadVerified": False, "phoneOpenVerified": False}
    review_path = root / "visual-review.json"
    if review_path.exists():
        review = read_receipt(root, "visual-review.json")
        require(review.get("artifacts") == result["artifacts"] and
                review.get("renderedPages") == result["renderedPages"] and
                review.get("status") == "passed", "visual review does not match current artifacts")
        result["visualReview"] = "passed"
    save(root / "verification.json", result)
    return {"nativeProcessing": "passed", "visualReview": result["visualReview"],
            "runDir": str(root), "ownerUploadVerified": False, "phoneOpenVerified": False}


def record_visual_review(args):
    verify(args)
    require(args.note and len(args.note.strip()) >= 10, "record a specific visual review note")
    result = read_receipt(args.run_dir, "verification.json")
    review = {"status": "passed", "reviewedAt": now(), "note": args.note,
              "artifacts": result["artifacts"], "renderedPages": result["renderedPages"]}
    save(args.run_dir / "visual-review.json", review)
    return verify(args)


def configured_owner(config):
    telegram = config.get("channels", {}).get("telegram", {})
    owners = telegram.get("allowFrom", [])
    require(telegram.get("enabled") is True and telegram.get("dmPolicy") == "allowlist" and
            len(owners) == 1 and isinstance(owners[0], str) and owners[0].isascii() and
            owners[0].isdigit() and int(owners[0]) > 0 and not telegram.get("accounts") and
            not telegram.get("direct"), "expected one configured owner without routing overrides")
    owner = owners[0]
    require(config.get("commands", {}).get("ownerAllowFrom") == ["telegram:" + owner],
            "command and Telegram owners differ")
    return owner


def transport_receipt(data, owner):
    require(isinstance(data, dict), "invalid Telegram result")
    payload = data.get("payload", {})
    require(data.get("action") == "send" and data.get("channel") == "telegram" and
            data.get("dryRun") is False and isinstance(payload, dict) and
            payload.get("ok") is True, "Telegram did not confirm an actual send")
    message_id = str(payload.get("messageId", ""))
    require(message_id.isascii() and message_id.isdigit() and int(message_id) > 0,
            "Telegram messageId missing")
    require(str(payload.get("chatId", "")) == owner, "Telegram receipt recipient mismatch")
    return {"action": "send", "channel": "telegram", "dryRun": False,
            "payload": {"ok": True, "messageId": message_id}, "recipientVerified": True}


def stage_media(root, name, digest):
    """Copy only checked synthetic bytes into the deployed native media allowlist."""
    run_key = hashlib.sha256(str(root).encode()).hexdigest()[:24]
    directory = STATE / "media" / "productivity-acceptance" / run_key
    for parent in reversed([directory, *directory.parents]):
        require(not parent.is_symlink(), "media staging parent is redirected")
        if parent == STATE or parent.is_relative_to(STATE):
            parent.mkdir(mode=0o700, exist_ok=True)
            require(parent.is_dir() and parent.stat().st_uid == os.getuid(),
                    "media staging directory has unexpected owner/type")
    source = safe_file(root, name)
    require(sha(source) == digest, "artifact changed before media staging")
    target = directory / (digest[:16] + "-" + source.name)
    require(not target.is_symlink(), "media staging target is redirected")
    if target.exists():
        require(target.is_file() and sha(target) == digest, "staged media content differs")
    else:
        fd = os.open(target, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "wb") as destination, source.open("rb") as origin:
            shutil.copyfileobj(origin, destination)
            destination.flush()
            os.fsync(destination.fileno())
        require(sha(target) == digest, "media staging hash mismatch")
    return target


def send(args):
    root = args.run_dir
    lock_path = root / ".send.lock"
    require(not lock_path.is_symlink(), "send lock is a symlink")
    with lock_path.open("a") as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise ValueError("another delivery run is active") from None
        verify(args)
        verification = read_receipt(root, "verification.json")
        require(verification["visualReview"] == "passed", "inspect all four rendered pages first")
        config_path = args.config.resolve(strict=True)
        # The managed wrapper pins this deployment and overrides environment
        # variables. Never infer a different live recipient from an unrelated file.
        require(config_path == (STATE / "openclaw.json").resolve(strict=True),
                "managed CLI pins the deployed config; alternate config cannot select a recipient")
        owner = configured_owner(json.loads(config_path.read_text(encoding="utf-8")))
        owner_hash = hashlib.sha256(owner.encode()).hexdigest()
        ledger = {"schemaVersion": 1, "recipientHash": owner_hash, "deliveries": {},
                  "scope": "synthetic-native-processing-and-telegram-outbound",
                  "ownerUploadVerified": False, "phoneOpenVerified": False}
        if (root / "delivery.json").exists():
            ledger = read_receipt(root, "delivery.json")
            require(ledger.get("recipientHash") == owner_hash, "configured recipient changed")
        for kind, name in ARTIFACTS:
            previous = ledger["deliveries"].get(name)
            digest = verification["artifacts"][name]
            require(sha(safe_file(root, name)) == digest, "artifact changed before delivery: " + name)
            if previous:
                require(previous.get("sha256") == digest, "previous delivery hash changed")
                require(previous.get("status") == "delivered",
                        "delivery outcome uncertain; inspect the saved receipt and Telegram before retrying")
                continue
            transport_artifact = subtitle_archive(root) if kind == "subtitles" else {
                "format": Path(name).suffix.lstrip("."), "path": name, "sha256": digest}
            media_path = stage_media(root, transport_artifact["path"], transport_artifact["sha256"])
            diagnostic_name = "diagnostics/" + kind + ("-zip" if kind == "subtitles" else "") + "-send.json"
            entry = {"type": kind, "sha256": digest, "checks": "passed",
                     "status": "sending", "startedAt": now(),
                     "transportArtifact": transport_artifact,
                     "stagedMedia": str(media_path.relative_to(STATE)),
                     "diagnostics": diagnostic_name}
            ledger["deliveries"][name] = entry
            save(root / "delivery.json", ledger)
            # Persist intent before sending; uncertainty is deliberately not auto-retried.
            try:
                env = os.environ.copy()
                env["OPENCLAW_CONFIG_PATH"] = str(config_path)
                env["OPENCLAW_STATE_DIR"] = str(config_path.parent)
                output = bounded([args.cli, "message", "send", "--channel", "telegram",
                                  "--target", owner, "--message",
                                  "OpenClaw 기능 점검용 합성 파일 · " +
                                  ("SRT 자막 ZIP (원본 speech-input.srt 포함)" if kind == "subtitles" else kind),
                                  "--media", media_path, "--force-document", "--silent", "--json"],
                                 timeout=120, env=env, diagnostic_path=root / diagnostic_name)
                entry["transport"] = transport_receipt(json.loads(output), owner)
                entry["status"] = "delivered"
                entry["deliveredAt"] = now()
            except (ValueError, OSError):
                entry["status"] = "uncertain"
                save(root / "delivery.json", ledger)
                raise
            save(root / "delivery.json", ledger)
        return {"nativeProcessing": "passed", "visualReview": "passed",
                "telegramOutbound": "passed", "deliveredArtifacts": len(ARTIFACTS),
                "messageIds": [entry["transport"]["payload"]["messageId"]
                               for entry in ledger["deliveries"].values()],
                "ownerUploadVerified": False, "phoneOpenVerified": False,
                "receipt": str(root / "delivery.json")}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=["prepare", "verify", "record-visual-review", "send"])
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--office", default="/opt/homebrew/bin/openclaw-office")
    parser.add_argument("--ffmpeg", default="/opt/homebrew/bin/ffmpeg")
    parser.add_argument("--whisper", default="/opt/homebrew/bin/whisper")
    parser.add_argument("--cli", type=Path, default=CLI)
    parser.add_argument("--config", type=Path, default=STATE / "openclaw.json")
    parser.add_argument("--note", help="specific review of all four displayed page images")
    args = parser.parse_args()
    require(not args.run_dir.is_symlink(), "run directory is a symlink")
    args.run_dir = args.run_dir.resolve()
    os.umask(0o077)
    action = {"prepare": prepare, "verify": verify, "record-visual-review": record_visual_review,
              "send": send}[args.mode]
    print(json.dumps(action(args), ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, AssertionError, ImportError) as error:
        # Provider diagnostics and credential-bearing arguments never enter stdout/stderr.
        print(json.dumps({"ok": False, "error": str(error) or type(error).__name__}), file=sys.stderr)
        sys.exit(1)
